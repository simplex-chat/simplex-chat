#!/usr/bin/env python3
import csv
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

WORKBOOK_FILE = "prices.xlsx"
USAGE = f"""Usage:
  python3 prices-workbook.py [-o <file>]
Reads apple-prices.csv, play-prices.csv, apple-price-points.csv, rates.csv and price-levels.csv from the current directory.
Writes {WORKBOOK_FILE} or <file> and refuses to overwrite an existing file. Requires openpyxl (pip install openpyxl)."""

FONT_NAME = "Arial"
INPUT_COLOR = "0000FF"
NOTE_COLOR = "595959"
UNUSED_COLOR = "A6A6A6"
HEADER_FILL = PatternFill("solid", fgColor="E7E6E6")
HIGHLIGHT_FILL = PatternFill("solid", fgColor="FCE4D6")

TIERS = [
    (1, 0.55, 7),
    (2, 0.40, 5),
    (3, 0, 3),
]
TOLERANCE = 0.035
STEPS = [
    (0, 1, 0.5, 0.5),
    (10, 1, 1, 1),
    *(
        row
        for unit in (1, 10, 100, 1000, 10000, 100000)
        for row in ((35 * unit, 5 * unit, 5 * unit, 5 * unit), (75 * unit, 10 * unit, 10 * unit, 10 * unit))
    ),
]
FAMILIES = [
    ("Supporter", 1, "yes"),
    ("Legend", 10, "no"),
]
ANNUAL_MONTHS = 6
TIER_OVERRIDES = {
    **dict.fromkeys(
        [
            "ALB", "AND", "AUT", "BEL", "BGR", "BIH", "CHE", "CYP", "CZE", "DEU", "DNK", "ESP", "EST", "FIN",
            "FRA", "FRO", "GBR", "GIB", "GRC", "HRV", "HUN", "IMN", "IRL", "ISL", "ITA", "LIE", "LTU", "LUX",
            "LVA", "MCO", "MKD", "MLT", "MNE", "NLD", "NOR", "POL", "PRT", "ROU", "SMR", "SRB", "SVK", "SVN",
            "SWE", "XKX",
        ],
        1,
    ),
    "RUS": 2,
    "CHN": 1,
}
AVOIDED_NUMBERS = [13, 66, 666]
AVOIDED_DIGITS = [
    ("CNY", 4, None),
    ("HKD", 4, None),
    ("MOP", 4, None),
    ("TWD", 4, None),
    ("JPY", 4, 9),
    ("KRW", 4, None),
    ("SGD", 4, None),
    ("VND", 4, None),
    ("MYR", 4, None),
]
MANUAL_PRICES = {
    ("ARS", 1): 10000,
    ("ETB", 1): 1150,
    ("HUF", 1): 2350,
    ("IDR", 1): 125000,
    ("IQD", 1): 10500,
    ("JPY", 1): 1050,
    ("KRW", 1): 10000,
    ("KZT", 3): 1350,
    ("MNT", 3): 10500,
    ("MXN", 1): 125,
    ("MYR", 3): 12.5,
    ("SGD", 1): 10,
    ("THB", 1): 235,
    ("TWD", 1): 215,
    ("VND", 3): 85000,
}
PRODUCTS = [
    ("apple", "SUBSCR_BADGE_SUPPORTER_MONTH_01", "", "Supporter monthly"),
    ("apple", "SUBSCR_BADGE_SUPPORTER_YEAR_01", "", "Supporter annual"),
    ("apple", "SUBSCR_BADGE_LEGEND_MONTH_01", "", "Legend monthly"),
    ("apple", "SUBSCR_BADGE_LEGEND_YEAR_01", "", "Legend annual"),
    ("apple", "BADGE_SUPPORTER_01", "", "Supporter monthly"),
    ("apple", "BADGE_LEGEND_01", "", "Legend monthly"),
    ("play", "subscr_badge_supporter_01", "subscr-badge-supporter-month-01", "Supporter monthly"),
    ("play", "subscr_badge_supporter_01", "subscr-badge-supporter-month-02", "Supporter monthly"),
    ("play", "subscr_badge_supporter_01", "subscr-badge-supporter-year-01", "Supporter annual"),
    ("play", "subscr_badge_legend_01", "subscr-badge-legend-month-01", "Legend monthly"),
    ("play", "subscr_badge_legend_01", "subscr-badge-legend-year-01", "Legend annual"),
    ("play", "badge_supporter_01", "", "Supporter monthly"),
    ("play", "badge_legend_01", "", "Legend monthly"),
]
APPLE_DOWN_LIMIT = 0.02
POINT_KEY_SCALE = 10**10
APPLE_MANUAL_PRICES = {
    ("RON", 1, "Supporter annual"): 200,
    ("TRY", 2, "Supporter annual"): 1500,
}
PLAY_MANUAL_PRICES = {
    ("SA", "subscr-badge-legend-year-01"): 1599,
}
APPLE_COLUMNS = ["type", "product_id", "territory", "currency", "current_price", "scheduled_price", "scheduled_date", "price"]
PLAY_COLUMNS = ["type", "product_id", "base_plan_id", "region", "currency", "current_price", "price"]
CHECK_COLUMNS = ["Country", "Tier", "Priced as", "Current, USD", "New, USD", "Change"]
COUNTRY_COLUMNS = [
    "Name",
    "ISO3",
    "ISO2",
    "Apple territory",
    "Play region",
    "Apple currency",
    "Play currency",
    "Price level",
    "Year",
    "Tier override",
    "Tier",
    "PPP, LCU per USD",
    "Atlas rate, LCU per USD",
    "Big Mac ratio",
    "Big Mac date",
    "Big Mac source",
]


def read_csv(path):
    with open(path, newline="", encoding="utf-8") as file:
        return list(csv.DictReader(file))


def number(text):
    return float(text) if text else None


def currency_by_code(rows, code_column):
    currencies = {}
    for row in rows:
        code, currency = row[code_column], row["currency"]
        if currencies.setdefault(code, currency) != currency:
            raise ValueError(f"{code} is priced in {currencies[code]} and {currency}")
    return currencies


def put(ws, row, column, value, number_format=None, color=None, bold=False, italic=False):
    cell = ws.cell(row, column, value)
    cell.font = Font(name=FONT_NAME, size=10, color=color, bold=bold, italic=italic)
    if number_format:
        cell.number_format = number_format
    return cell


def put_header(ws, row, headers, first_column=1):
    for column, header in enumerate(headers, start=first_column):
        put(ws, row, column, header, bold=True).fill = HEADER_FILL


def put_note(ws, row, text):
    put(ws, row, 1, text, color=NOTE_COLOR, italic=True)


def set_widths(ws, widths):
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def finish_table(ws, columns, last_row):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(columns)}{last_row}"


def avoided_formula(settings, price, currency):
    numbers = settings["avoided_numbers"]
    digits = [f'IFERROR(INDEX({column},MATCH({currency},{settings["avoided_currencies"]},0))&"","")' for column in (settings["avoided_first"], settings["avoided_second"])]
    contains_digit = ",".join(f'AND({digit}<>"",ISNUMBER(SEARCH({digit},{price}&"")))' for digit in digits)
    return (
        f"OR({contains_digit},"
        f"COUNTIF({numbers},{price}/10^(INT(LOG10({price}))-1))+COUNTIF({numbers},{price}/10^(INT(LOG10({price}))-2))"
        f"+COUNTIF({numbers},MOD(INT({price}/1000),1000))+COUNTIF({numbers},MOD(INT({price}/10^6),1000))>0)"
    )


def add_validation(ws, cells, error, **rule):
    validation = DataValidation(showErrorMessage=True, errorStyle="stop", error=error, **rule)
    validation.add(cells)
    ws.add_data_validation(validation)


def build_settings(ws, rate_date):
    put(ws, 1, 1, "SimpleX badge prices", bold=True)
    put_note(ws, 2, "Blue cells are inputs. Every price is computed from them.")

    put(ws, 4, 1, "Tiers", bold=True)
    put_header(ws, 5, ["Tier", "Price level from", "Supporter monthly, USD"])
    tier_row = 6
    for offset, (tier, level, usd) in enumerate(TIERS):
        row = tier_row + offset
        put(ws, row, 1, tier)
        put(ws, row, 2, level, "0.00", color=INPUT_COLOR if tier < len(TIERS) else None)
        put(ws, row, 3, usd, color=INPUT_COLOR)
    put_note(ws, tier_row + len(TIERS), "Price level: World Bank PPP factor ÷ Atlas rate, US = 1. Countries without a price level are in tier 1. A tier entered in Countries column J replaces the computed tier: 1 for European countries outside the former USSR, the Baltic states and China, 2 for Russia.")

    put(ws, 11, 1, "Rounding", bold=True)
    tolerance_row = 12
    put(ws, tolerance_row, 1, "Tolerance")
    put(ws, tolerance_row, 2, TOLERANCE, "0.0%", color=INPUT_COLOR)
    put_note(ws, 13, "The exact amount is rounded down to the step when the result is within the tolerance, otherwise up.")
    put_header(ws, 14, ["Amount from", "Step, tier 1", "Step, tier 2", "Step, tier 3"])
    step_row = 15
    for offset, values in enumerate(STEPS):
        for column, value in enumerate(values, start=1):
            put(ws, step_row + offset, column, value, color=INPUT_COLOR)
    step_last = step_row + len(STEPS) - 1
    put_note(ws, step_last + 1, "Each step is used from its amount up to the next row's amount.")

    family_title = step_last + 3
    put(ws, family_title, 1, "Families", bold=True)
    put_header(ws, family_title + 1, ["Family", "× supporter monthly", "Discounted"])
    family_row = family_title + 2
    for offset, (name, multiple, discounted) in enumerate(FAMILIES):
        row = family_row + offset
        put(ws, row, 1, name)
        put(ws, row, 2, multiple, color=INPUT_COLOR)
        put(ws, row, 3, discounted, color=INPUT_COLOR)
    family_last = family_row + len(FAMILIES) - 1
    put_note(ws, family_last + 1, "Monthly price = multiple × the final supporter monthly price of the country's tier when discounted, of tier 1 otherwise.")
    add_validation(ws, f"C{family_row}:C{family_last}", "Enter yes or no.", type="list", formula1='"yes,no"', allow_blank=False)
    annual_row = family_last + 2
    put(ws, annual_row, 1, "Annual months")
    put(ws, annual_row, 2, ANNUAL_MONTHS, color=INPUT_COLOR)
    put_note(ws, annual_row + 1, "Annual price = annual months × monthly price.")
    add_validation(ws, f"B{annual_row}", "Enter a whole number of months.", type="whole", operator="between", formula1="1", formula2="12", allow_blank=False)

    avoided_title = annual_row + 3
    put(ws, avoided_title, 1, "Avoided numbers", bold=True)
    numbers_row = avoided_title + 1
    put(ws, numbers_row, 1, "Everywhere")
    for column, avoided in enumerate(AVOIDED_NUMBERS, start=2):
        put(ws, numbers_row, column, avoided, color=INPUT_COLOR)
    put_note(ws, numbers_row + 1, "A price is flagged in the Warning columns of the Apple and Play tabs, and shaded orange on the Prices tab in Excel, when its significant digits, or its thousands or millions, equal one of these numbers.")
    put_header(ws, numbers_row + 2, ["Currency", "Digit", "Second digit"])
    digit_row = numbers_row + 3
    for offset, (currency, digit, second_digit) in enumerate(AVOIDED_DIGITS):
        put(ws, digit_row + offset, 1, currency, color=INPUT_COLOR)
        put(ws, digit_row + offset, 2, digit, color=INPUT_COLOR)
        put(ws, digit_row + offset, 3, second_digit, color=INPUT_COLOR)
    digit_last = digit_row + len(AVOIDED_DIGITS) - 1
    put_note(ws, digit_last + 1, "A price in the currency is flagged when it contains either digit. A flagged price is replaced by a value in the Manual column of the Prices tab.")

    product_title = digit_last + 3
    put(ws, product_title, 1, "Store products", bold=True)
    put_header(ws, product_title + 1, ["Store", "Product ID", "Base plan ID", "Priced as", "Key"])
    product_row = product_title + 2
    for offset, (store, product_id, base_plan_id, priced_as) in enumerate(PRODUCTS):
        row = product_row + offset
        put(ws, row, 1, store)
        put(ws, row, 2, product_id)
        put(ws, row, 3, base_plan_id or None)
        put(ws, row, 4, priced_as, color=INPUT_COLOR)
        put(ws, row, 5, f'=A{row}&"|"&B{row}&"|"&C{row}')
    product_last = product_row + len(PRODUCTS) - 1

    points_title = product_last + 2
    put(ws, points_title, 1, "Apple price points", bold=True)
    apple_limit_row = points_title + 1
    put(ws, apple_limit_row, 1, "Round-down limit")
    put(ws, apple_limit_row, 2, APPLE_DOWN_LIMIT, "0.0%", color=INPUT_COLOR)
    put_note(ws, apple_limit_row + 1, "Apple prices are calculated on the Apple prices tab, one row per currency and tier. A price is the target when Apple offers it; otherwise the highest round price point at or below the target when that is within the limit, else the lowest round price point above it.")
    put_note(ws, apple_limit_row + 2, "Round: a whole number, or a multiple of 0.5 below 10; the Apple points tab lists the round points and the others. Target: the workbook price; for annual prices, annual months × the Apple monthly price of the same row.")
    put_note(ws, apple_limit_row + 3, "A value in the Manual columns of the Apple prices tab replaces the Apple price, and one in the Play manual column of the Play tab replaces the Play price. Points: apple-price-points.csv.")

    export_title = apple_limit_row + 5
    put(ws, export_title, 1, "Import", bold=True)
    put_note(ws, export_title + 1, "Save sheet apple-prices.csv as CSV UTF-8 to scripts/prices/apple-prices.csv, then run: node apple-prices.js import <issuer_id> <path/to/AuthKey_XXXX.p8> --dry-run")
    put_note(ws, export_title + 2, "Save sheet play-prices.csv as CSV UTF-8 to scripts/prices/play-prices.csv, then run: node play-prices.js import <service-account.json> --dry-run")
    put_note(ws, export_title + 3, "Both import scripts read a decimal point (3.5), so save from a region that writes 3.5, not 3,5.")
    put_note(ws, export_title + 4, f"Rates: rates.csv, Frankfurter, {rate_date}. Price levels: price-levels.csv.")

    set_widths(ws, [22, 38, 34, 20, 44])
    return {
        "tier1_level": f"Settings!$B${tier_row}",
        "tier2_level": f"Settings!$B${tier_row + 1}",
        "tier_usd": f"Settings!$C${tier_row}:$C${tier_row + len(TIERS) - 1}",
        "tolerance": f"Settings!$B${tolerance_row}",
        "step_from": f"Settings!$A${step_row}:$A${step_last}",
        "steps": f"Settings!$B${step_row}:$D${step_last}",
        "family_rows": [family_row + offset for offset in range(len(FAMILIES))],
        "annual_months": f"Settings!$B${annual_row}",
        "avoided_numbers": f"Settings!$B${numbers_row}:${get_column_letter(1 + len(AVOIDED_NUMBERS))}${numbers_row}",
        "avoided_currencies": f"Settings!$A${digit_row}:$A${digit_last}",
        "avoided_first": f"Settings!$B${digit_row}:$B${digit_last}",
        "avoided_second": f"Settings!$C${digit_row}:$C${digit_last}",
        "product_cells": f"D{product_row}:D{product_last}",
        "product_priced_as": f"Settings!$D${product_row}:$D${product_last}",
        "product_key": f"Settings!$E${product_row}:$E${product_last}",
        "apple_down_limit": f"Settings!$B${apple_limit_row}",
    }


def country_rows(levels, apple_currency, play_currency):
    known_apple = {row["apple_territory"] for row in levels}
    known_play = {row["google_region"] for row in levels}
    added = [
        *({"name": code, "apple_territory": code} for code in sorted(apple_currency.keys() - known_apple)),
        *({"name": code, "google_region": code} for code in sorted(play_currency.keys() - known_play)),
    ]
    return sorted([*levels, *added], key=lambda row: row["name"])


def build_countries(ws, countries, apple_currency, play_currency, settings):
    put_header(ws, 1, COUNTRY_COLUMNS)
    last_row = len(countries) + 1
    for row, country in enumerate(countries, start=2):
        apple = country.get("apple_territory", "")
        play = country.get("google_region", "")
        year = country.get("price_level_year", "")
        put(ws, row, 1, country["name"])
        put(ws, row, 2, country.get("iso3") or None)
        put(ws, row, 3, country.get("iso2") or None)
        put(ws, row, 4, apple or None)
        put(ws, row, 5, play or None)
        put(ws, row, 6, apple_currency.get(apple))
        put(ws, row, 7, play_currency.get(play))
        put(ws, row, 8, number(country.get("price_level", "")), "0.0000")
        put(ws, row, 9, int(year) if year else None)
        put(ws, row, 10, TIER_OVERRIDES.get(country.get("iso3")), color=INPUT_COLOR)
        put(ws, row, 11, f'=IF(J{row}<>"",J{row},IF(H{row}="",1,IF(H{row}>={settings["tier1_level"]},1,IF(H{row}>={settings["tier2_level"]},2,3))))')
        put(ws, row, 12, number(country.get("ppp_lcu_per_usd", "")))
        put(ws, row, 13, number(country.get("atls_lcu_per_usd", "")))
        put(ws, row, 14, number(country.get("big_mac_ratio", "")), "0.000")
        put(ws, row, 15, country.get("big_mac_date") or None)
        put(ws, row, 16, country.get("big_mac_source") or None)
    add_validation(ws, f"J2:J{last_row}", "Enter tier 1, 2 or 3, or leave empty.", type="whole", operator="between", formula1="1", formula2=str(len(TIERS)), allow_blank=True)
    finish_table(ws, len(COUNTRY_COLUMNS), last_row)
    set_widths(ws, [30, 7, 7, 10, 9, 10, 9, 11, 7, 10, 6, 14, 14, 11, 12, 11])
    return {
        "name": f"Countries!$A$2:$A${last_row}",
        "apple_code": f"Countries!$D$2:$D${last_row}",
        "play_code": f"Countries!$E$2:$E${last_row}",
        "apple_currency": f"Countries!$F$2:$F${last_row}",
        "play_currency": f"Countries!$G$2:$G${last_row}",
        "tier": f"Countries!$K$2:$K${last_row}",
    }


def build_rates(ws, rates):
    put_header(ws, 1, ["Currency", "Units per USD", "Date", "Source"])
    last_row = len(rates) + 1
    for row, rate in enumerate(rates, start=2):
        put(ws, row, 1, rate["currency"])
        put(ws, row, 2, float(rate["units_per_usd"]))
        put(ws, row, 3, rate["rate_date"])
        put(ws, row, 4, rate["source"])
    finish_table(ws, 4, last_row)
    set_widths(ws, [10, 14, 12, 40])
    return {
        "currency": f"Rates!$A$2:$A${last_row}",
        "units": f"Rates!$B$2:$B${last_row}",
    }


def build_prices(ws, currencies, settings, countries, rates):
    headers = ["Key", "Currency", "Tier", "Apple territories", "Play regions", "Units per USD", "Exact", "Step", "Down", "Down vs exact", "Rounded", "Manual", "Final", "Final vs exact"]
    first_price_column = len(headers) + 1
    last_price_column = len(headers) + 2 * len(FAMILIES)
    put_header(ws, 1, headers)
    for offset, family_row in enumerate(settings["family_rows"]):
        monthly_column = first_price_column + 2 * offset
        put(ws, 1, monthly_column, f'=Settings!$A${family_row}&" monthly"', bold=True).fill = HEADER_FILL
        put(ws, 1, monthly_column + 1, f'=Settings!$A${family_row}&" annual"', bold=True).fill = HEADER_FILL
    pairs = [(currency, tier) for currency in currencies for tier, _, _ in TIERS]
    last_row = len(pairs) + 1
    keys = f"$A$2:$A${last_row}"
    final = f"$M$2:$M${last_row}"
    for row, (currency, tier) in enumerate(pairs, start=2):
        tier1_final = f'INDEX({final},MATCH($B{row}&"|1",{keys},0))'
        put(ws, row, 1, f'=B{row}&"|"&C{row}', color=NOTE_COLOR)
        put(ws, row, 2, currency)
        put(ws, row, 3, tier)
        put(ws, row, 4, f"=COUNTIFS({countries['apple_currency']},$B{row},{countries['tier']},$C{row})")
        put(ws, row, 5, f"=COUNTIFS({countries['play_currency']},$B{row},{countries['tier']},$C{row})")
        put(ws, row, 6, f"=INDEX({rates['units']},MATCH($B{row},{rates['currency']},0))")
        put(ws, row, 7, f"=INDEX({settings['tier_usd']},$C{row})*F{row}", "#,##0.00")
        put(ws, row, 8, f"=INDEX({settings['steps']},MATCH(G{row},{settings['step_from']},1),$C{row})")
        put(ws, row, 9, f"=FLOOR(G{row},H{row})")
        put(ws, row, 10, f"=I{row}/G{row}-1", "0.0%")
        put(ws, row, 11, f"=IF(I{row}>=G{row}*(1-{settings['tolerance']}),I{row},CEILING(G{row},H{row}))")
        put(ws, row, 12, MANUAL_PRICES.get((currency, tier)), color=INPUT_COLOR)
        put(ws, row, 13, f'=IF(L{row}="",K{row},L{row})')
        put(ws, row, 14, f"=M{row}/G{row}-1", "0.0%")
        for offset, family_row in enumerate(settings["family_rows"]):
            monthly_column = first_price_column + 2 * offset
            monthly = f"{get_column_letter(monthly_column)}{row}"
            put(ws, row, monthly_column, f'=Settings!$B${family_row}*IF(Settings!$C${family_row}="yes",$M{row},{tier1_final})')
            put(ws, row, monthly_column + 1, f"={settings['annual_months']}*{monthly}")
    first_price, last_price = get_column_letter(first_price_column), get_column_letter(last_price_column)
    add_validation(ws, f"L2:L{last_row}", "Enter a positive price or leave empty.", type="decimal", operator="greaterThan", formula1="0", allow_blank=True)
    ws.conditional_formatting.add(f"A2:{last_price}{last_row}", FormulaRule(formula=["AND($C2<>1,$D2+$E2=0)"], font=Font(color=UNUSED_COLOR)))
    ws.conditional_formatting.add(f"M2:M{last_row} {first_price}2:{last_price}{last_row}", FormulaRule(formula=[avoided_formula(settings, "M2", "$B2")], fill=HIGHLIGHT_FILL))
    finish_table(ws, last_price_column, last_row)
    set_widths(ws, [9, 9, 5, 10, 8, 12, 12, 8, 10, 10, 10, 10, 10, 10, 18, 18, 18, 18])
    put_note(ws, last_row + 2, "Exact: the tier's supporter monthly USD price at the rate. Down: the exact amount rounded down to the step. Rounded: Down when it is within the tolerance, otherwise rounded up.")
    put_note(ws, last_row + 3, "Final: Manual when entered, otherwise Rounded. In Excel, orange cells contain an avoided number from Settings; Numbers shows these only in the Warning columns of the Apple and Play tabs.")
    put_note(ws, last_row + 4, "Grey rows: tiers 2 and 3 without territories. Tier-1 rows also set the legend prices of every tier.")
    return {
        "key": f"Prices!{keys}",
        "names": f"Prices!${first_price}$1:${last_price}$1",
        "prices": f"Prices!${first_price}$2:${last_price}${last_row}",
    }


def build_store(ws, columns, rows, code_column, country_codes, product_key, settings, countries, rates, price, extra):
    headers = [*columns, *CHECK_COLUMNS, *(header for header, _, _, _ in extra)]
    put_header(ws, 1, headers)
    letter = {name: get_column_letter(index) for index, name in enumerate(headers, start=1)}
    code, currency, current, price_column = (letter[name] for name in (code_column, "currency", "current_price", "price"))
    last_row = len(rows) + 1
    for row, data in enumerate(rows, start=2):
        for index, name in enumerate(columns, start=1):
            put(ws, row, index, price(row, letter) if name == "price" else data[name] or None)
        rate = f"INDEX({rates['units']},MATCH(${currency}{row},{rates['currency']},0))"
        put(ws, row, len(columns) + 1, f"=INDEX({countries['name']},MATCH(${code}{row},{country_codes},0))")
        put(ws, row, len(columns) + 2, f"=INDEX({countries['tier']},MATCH(${code}{row},{country_codes},0))")
        put(ws, row, len(columns) + 3, f"=INDEX({settings['product_priced_as']},MATCH({product_key(row)},{settings['product_key']},0))")
        put(ws, row, len(columns) + 4, f"=VALUE({current}{row})/{rate}", "#,##0.00")
        put(ws, row, len(columns) + 5, f"={price_column}{row}/{rate}", "#,##0.00")
        put(ws, row, len(columns) + 6, f"={price_column}{row}/VALUE({current}{row})-1", "0.0%")
        for offset, (_, value, number_format, color) in enumerate(extra, start=len(columns) + len(CHECK_COLUMNS) + 1):
            put(ws, row, offset, value(row, letter), number_format, color=color)
    finish_table(ws, len(headers), last_row)
    return last_row, letter


def workbook_price(prices, row, letter):
    return f'=INDEX({prices["prices"]},MATCH(${letter["currency"]}{row}&"|"&${letter["Tier"]}{row},{prices["key"]},0),MATCH(${letter["Priced as"]}{row},{prices["names"]},0))'


def points_by_currency(rows):
    point_sets = defaultdict(set)
    for row in rows:
        point_sets[(row["type"], row["territory"], row["currency"])].add(float(row["customer_price"]))
    by_currency = {}
    for (kind, territory, currency), points in point_sets.items():
        if by_currency.setdefault(currency, points) != points:
            raise ValueError(f"{kind} price points of {territory} differ from other {currency} territories")
    return by_currency


def is_round(price):
    return price == int(price) or (price < 10 and price * 2 == int(price * 2))


def put_rows(ws, first_column, rows, number_formats):
    for row, values in enumerate(rows, start=2):
        for column, (value, number_format) in enumerate(zip(values, number_formats), start=first_column):
            put(ws, row, column, value, number_format)


def build_points(ws, by_currency):
    currencies = sorted(by_currency)
    points = [(currency, point, (index + 1) * POINT_KEY_SCALE + point) for index, currency in enumerate(currencies) for point in sorted(by_currency[currency])]
    round_points = [point for point in points if is_round(point[1])]
    other_points = [point for point in points if not is_round(point[1])]
    put_header(ws, 1, ["Round currency", "Round price", "Round key"])
    put_header(ws, 1, ["Other currency", "Other price", "Other key"], first_column=5)
    put_header(ws, 1, ["Currencies"], first_column=9)
    put_rows(ws, 1, round_points, [None, None, "0.00"])
    put_rows(ws, 5, other_points, [None, None, "0.00"])
    put_rows(ws, 9, [(currency,) for currency in currencies], [None])
    finish_table(ws, 3, len(round_points) + 1)
    set_widths(ws, [15, 14, 18, 3, 15, 14, 18, 3, 12])
    put_note(ws, len(round_points) + 3, "Key: the currency's position in Currencies × 10^10 + price; both lists are sorted by key.")
    return {
        "currencies": f"'Apple points'!$I$2:$I${len(currencies) + 1}",
        "round_currency": f"'Apple points'!$A$2:$A${len(round_points) + 1}",
        "round_price": f"'Apple points'!$B$2:$B${len(round_points) + 1}",
        "round_keys": f"'Apple points'!$C$2:$C${len(round_points) + 1}",
        "other_keys": f"'Apple points'!$G$2:$G${len(other_points) + 1}",
    }


def offered_formula(keys, currency_key, price):
    return f"IFERROR(INDEX({keys},MATCH({currency_key}+{price},{keys},1))={currency_key}+{price},FALSE)"


def apple_price_formula(settings, points, currency, currency_key, target, found_row, manual):
    below = f"IF({found_row}=0,0,IF(INDEX({points['round_currency']},{found_row})={currency},INDEX({points['round_price']},{found_row}),0))"
    above = f"IFERROR(IF(INDEX({points['round_currency']},{found_row}+1)={currency},INDEX({points['round_price']},{found_row}+1),0),0)"
    round_target = f"OR({target}=INT({target}),AND({target}<10,{target}*2=INT({target}*2)))"
    offered = f"IF({round_target},{below}={target},{offered_formula(points['other_keys'], currency_key, target)})"
    nearby = f"IF(OR({below}>={target}*(1-{settings['apple_down_limit']}),{above}=0),{below},{above})"
    return f'=IF({manual}<>"",{manual},IF({offered},{target},{nearby}))'


def apple_warning_formula(settings, points, currency, currency_key, price, manual):
    offered = f"OR({offered_formula(points['round_keys'], currency_key, price)},{offered_formula(points['other_keys'], currency_key, price)})"
    avoided = avoided_formula(settings, price, currency)
    return f'=IF({price}=0,"no round price point",IF(IF({manual}<>"",NOT({offered}),FALSE),"not an Apple price point",IF({avoided},"avoided number","")))'


def build_apple_prices(ws, currencies, settings, prices, points, countries):
    price_count = 2 * len(FAMILIES)
    price_names = [f"{name} {period}" for name, _, _ in FAMILIES for period in ("monthly", "annual")]
    headers = ["Key", "Currency", "Tier", "Apple territories"]
    blocks = ["Apple", "Warning", "Manual", "Workbook", "Target", "Row"]
    block_start = {block: len(headers) + 1 + offset * price_count for offset, block in enumerate(blocks)}
    key_column = len(headers) + len(blocks) * price_count + 1
    put_header(ws, 1, headers)
    for block, start in block_start.items():
        for index in range(price_count):
            name = f"INDEX({prices['names']},{index + 1})"
            put(ws, 1, start + index, f"={name}" if block == "Apple" else f'="{block} "&{name}', bold=True).fill = HEADER_FILL
    put(ws, 1, key_column, "Currency key", bold=True).fill = HEADER_FILL
    pairs = [(currency, tier) for currency in currencies for tier, _, _ in TIERS]
    last_row = len(pairs) + 1

    def cell(block, index, row):
        return f"${get_column_letter(block_start[block] + index)}{row}"

    for row, (currency, tier) in enumerate(pairs, start=2):
        currency_key = f"${get_column_letter(key_column)}{row}"
        put(ws, row, 1, f'=B{row}&"|"&C{row}', color=NOTE_COLOR)
        put(ws, row, 2, currency)
        put(ws, row, 3, tier)
        put(ws, row, 4, f"=COUNTIFS({countries['apple_currency']},$B{row},{countries['tier']},$C{row})")
        for index, name in enumerate(price_names):
            target, found_row, manual, apple = (cell(block, index, row) for block in ("Target", "Row", "Manual", "Apple"))
            target_formula = f"={cell('Workbook', index, row)}" if index % 2 == 0 else f"={settings['annual_months']}*{cell('Apple', index - 1, row)}"
            put(ws, row, block_start["Workbook"] + index, f"=INDEX({prices['prices']},MATCH($A{row},{prices['key']},0),{index + 1})")
            put(ws, row, block_start["Target"] + index, target_formula)
            put(ws, row, block_start["Row"] + index, f"=IFERROR(MATCH({currency_key}+{target},{points['round_keys']},1),0)")
            put(ws, row, block_start["Manual"] + index, APPLE_MANUAL_PRICES.get((currency, tier, name)), color=INPUT_COLOR)
            put(ws, row, block_start["Apple"] + index, apple_price_formula(settings, points, f"$B{row}", currency_key, target, found_row, manual))
            put(ws, row, block_start["Warning"] + index, apple_warning_formula(settings, points, f"$B{row}", currency_key, apple, manual))
        put(ws, row, key_column, f"=MATCH($B{row},{points['currencies']},0)*{POINT_KEY_SCALE}", "0")
    first_apple, last_apple = get_column_letter(block_start["Apple"]), get_column_letter(block_start["Apple"] + price_count - 1)
    first_warning, last_warning = get_column_letter(block_start["Warning"]), get_column_letter(block_start["Warning"] + price_count - 1)
    first_manual, last_manual = get_column_letter(block_start["Manual"]), get_column_letter(block_start["Manual"] + price_count - 1)
    add_validation(ws, f"{first_manual}2:{last_manual}{last_row}", "Enter a positive Apple price point or leave empty.", type="decimal", operator="greaterThan", formula1="0", allow_blank=True)
    ws.conditional_formatting.add(f"A2:{get_column_letter(key_column)}{last_row}", FormulaRule(formula=["$D2=0"], font=Font(color=UNUSED_COLOR)))
    ws.conditional_formatting.add(f"{first_apple}2:{last_apple}{last_row}", FormulaRule(formula=[f'{first_warning}2<>""'], fill=HIGHLIGHT_FILL))
    finish_table(ws, key_column, last_row)
    set_widths(ws, [9, 9, 5, 10, *[16] * (len(blocks) * price_count), 16])
    put_note(ws, last_row + 2, "One row per Apple currency and tier. Apple: the price sent to Apple. Manual: replaces the Apple price. Workbook: the price from the Prices tab.")
    put_note(ws, last_row + 3, "Target: the workbook price; for annual prices, annual months × the Apple monthly price. Row: the position of the highest round price point at or below the target on the Apple points tab.")
    put_note(ws, last_row + 4, "Grey rows: no Apple territory uses this currency in this tier.")
    return {
        "key": f"'Apple prices'!$A$2:$A${last_row}",
        "names": f"'Apple prices'!${first_apple}$1:${last_apple}$1",
        "prices": f"'Apple prices'!${first_apple}$2:${last_apple}${last_row}",
        "warnings": f"'Apple prices'!${first_warning}$2:${last_warning}${last_row}",
    }


def apple_columns(prices, apple_prices):
    def position(row, letter):
        return f'MATCH(${letter["currency"]}{row}&"|"&${letter["Tier"]}{row},{apple_prices["key"]},0),MATCH(${letter["Priced as"]}{row},{apple_prices["names"]},0)'

    def apple_price(row, letter):
        return f"=INDEX({apple_prices['prices']},{position(row, letter)})"

    extra = [
        ("Workbook price", lambda row, letter: workbook_price(prices, row, letter), None, None),
        ("Apple vs workbook", lambda row, letter: f"=${letter['price']}{row}/${letter['Workbook price']}{row}-1", "0.0%", None),
        ("Warning", lambda row, letter: f'=INDEX({apple_prices["warnings"]},{position(row, letter)})&""', None, None),
    ]
    return apple_price, extra


def play_columns(settings, prices, play_rows):
    rows_by_number = dict(enumerate(play_rows, start=2))

    def play_price(row, letter):
        manual, workbook = f"${letter['Play manual']}{row}", f"${letter['Workbook price']}{row}"
        return f'=IF({manual}<>"",{manual},{workbook})'

    def play_manual(row, letter):
        data = rows_by_number[row]
        return PLAY_MANUAL_PRICES.get((data["region"], data["base_plan_id"] or data["product_id"]))

    def warning(row, letter):
        return f'=IF({avoided_formula(settings, f"${letter["price"]}{row}", f"$E{row}")},"avoided number","")'

    extra = [
        ("Workbook price", lambda row, letter: workbook_price(prices, row, letter), None, None),
        ("Play manual", play_manual, None, INPUT_COLOR),
        ("Warning", warning, None, None),
    ]
    return play_price, extra


def build_export(ws, source, columns, last_row, price_formula):
    put_header(ws, 1, columns)
    for row in range(2, last_row + 1):
        for index, name in enumerate(columns, start=1):
            cell = f"{source}!{get_column_letter(index)}{row}"
            put(ws, row, index, price_formula(row) if name == "price" else f'=IF({cell}="","",{cell})')


def output_path(args):
    if not args:
        return Path(WORKBOOK_FILE)
    if len(args) == 2 and args[0] == "-o":
        return Path(args[1])
    sys.exit(USAGE)


def main():
    output = output_path(sys.argv[1:])
    if output.exists():
        sys.exit(f"{output} already exists, nothing was written")
    apple_rows = read_csv("apple-prices.csv")
    play_rows = read_csv("play-prices.csv")
    rates = read_csv("rates.csv")
    levels = read_csv("price-levels.csv")
    points = points_by_currency(read_csv("apple-price-points.csv"))
    apple_currency = currency_by_code(apple_rows, "territory")
    play_currency = currency_by_code(play_rows, "region")
    currencies = sorted({*apple_currency.values(), *play_currency.values()})

    wb = Workbook()
    settings_ws = wb.active
    settings_ws.title = "Settings"
    prices_ws = wb.create_sheet("Prices")
    apple_prices_ws = wb.create_sheet("Apple prices")
    countries_ws = wb.create_sheet("Countries")
    apple_ws = wb.create_sheet("Apple")
    play_ws = wb.create_sheet("Play")
    rates_ws = wb.create_sheet("Rates")
    points_ws = wb.create_sheet("Apple points")
    apple_export_ws = wb.create_sheet("apple-prices.csv")
    play_export_ws = wb.create_sheet("play-prices.csv")

    settings = build_settings(settings_ws, max(rate["rate_date"] for rate in rates))
    countries = build_countries(countries_ws, country_rows(levels, apple_currency, play_currency), apple_currency, play_currency, settings)
    rate_ranges = build_rates(rates_ws, rates)
    prices = build_prices(prices_ws, currencies, settings, countries, rate_ranges)
    add_validation(settings_ws, settings["product_cells"], "Choose a price column of the Prices sheet.", type="list", formula1=prices["names"], allow_blank=False)
    point_ranges = build_points(points_ws, points)
    apple_prices = build_apple_prices(apple_prices_ws, sorted(set(apple_currency.values())), settings, prices, point_ranges, countries)
    apple_price, apple_extra = apple_columns(prices, apple_prices)
    apple_last, apple_letter = build_store(apple_ws, APPLE_COLUMNS, apple_rows, "territory", countries["apple_code"], lambda row: f'"apple|"&$B{row}&"|"', settings, countries, rate_ranges, apple_price, apple_extra)
    play_price, play_extra = play_columns(settings, prices, play_rows)
    play_last, play_letter = build_store(play_ws, PLAY_COLUMNS, play_rows, "region", countries["play_code"], lambda row: f'"play|"&$B{row}&"|"&$C{row}', settings, countries, rate_ranges, play_price, play_extra)
    play_manual_column = play_letter["Play manual"]
    add_validation(play_ws, f"{play_manual_column}2:{play_manual_column}{play_last}", "Enter a positive price or leave empty.", type="decimal", operator="greaterThan", formula1="0", allow_blank=True)
    for ws, letter, last_row in ((apple_ws, apple_letter, apple_last), (play_ws, play_letter, play_last)):
        ws.conditional_formatting.add(f"{letter['price']}2:{letter['price']}{last_row}", FormulaRule(formula=[f'${letter["Warning"]}2<>""'], fill=HIGHLIGHT_FILL))
    set_widths(apple_ws, [12, 36, 9, 9, 12, 14, 14, 10, 30, 6, 18, 12, 10, 9, 14, 12, 24])
    set_widths(play_ws, [12, 26, 32, 7, 9, 12, 10, 30, 6, 18, 12, 10, 9, 14, 12, 18])
    apple_decimals = 'IF(ROUND(Apple!H{row}*10,6)=ROUND(Apple!H{row}*10,0),1,2)'
    build_export(apple_export_ws, "Apple", APPLE_COLUMNS, apple_last, lambda row: f'=FIXED(Apple!H{row},IF(ISNUMBER(SEARCH(".",Apple!E{row})),{apple_decimals.format(row=row)},0),TRUE)')
    play_decimals = 'IF(ROUND(Play!G{row},6)=ROUND(Play!G{row},0),0,IF(ROUND(Play!G{row}*10,6)=ROUND(Play!G{row}*10,0),1,2))'
    build_export(play_export_ws, "Play", PLAY_COLUMNS, play_last, lambda row: f"=FIXED(Play!G{row},{play_decimals.format(row=row)},TRUE)")

    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(output)
    print(f"Wrote {output}: {len(currencies)} currencies, {len(apple_rows)} Apple rows, {len(play_rows)} Play rows")


main()
